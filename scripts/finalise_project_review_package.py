"""Validate the evidence collection and package only the intended deliverables."""
import csv
import hashlib
import html
import json
from pathlib import Path
import re
import zipfile
from docx import Document
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
PROJECT=ROOT/'20106UG03'
REC=PROJECT/'Review_Records'
OUTPUT=ROOT/'outputs'
OUTPUT.mkdir(exist_ok=True)
register=json.loads((REC/'Paper_Register.json').read_text(encoding='utf-8'))
summaries=json.loads((REC/'source_checked_summaries.json').read_text(encoding='utf-8'))
cloud=json.loads((REC/'Drive_PDF_Readback.json').read_text(encoding='utf-8'))
remote={x['name']:x for x in cloud['files']}
checks=[]
assert len(register)==len(summaries)==101
assert len({r['doi'] for r in register})==101
assert len({r['sha256'] for r in register})==101
for r in register:
    p=PROJECT/r['file'];assert p.is_file()
    assert hashlib.sha256(p.read_bytes()).hexdigest()==r['sha256']
    assert p.name in remote
    assert int(remote[p.name]['size'])==p.stat().st_size
    assert remote[p.name]['shared'] is False
    for field in ('title','journal','authors'):
        r[field]=html.unescape(r[field])
    checks.append({'review_id':r['review_id'],'pdf_hash_verified':True,
                   'drive_name_and_size_verified':True})
(REC/'Paper_Register.json').write_text(json.dumps(register,ensure_ascii=False,indent=2),encoding='utf-8')
with (PROJECT/'Paper_Register.csv').open('w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(register[0]));w.writeheader();w.writerows(register)
docs=[]
for p in sorted(PROJECT.glob('*.docx')):
    d=Document(p)
    assert not any(d.styles['Title'].element.iter(qn('w:pBdr')))
    text='\n'.join(x.text for x in d.paragraphs)
    assert not re.search(r'\[@|&amp;|turn\d+(?:view|search)',text)
    docs.append({'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),
                 'paragraphs':len(d.paragraphs),'tables':len(d.tables)})
assert len(docs)==2
audit={'date':'2026-10-31','project_number':'20106UG03','studies':101,
       'distinct_pdf_hashes':101,'summaries':101,'drive_pdfs':len(remote),
       'drive_name_and_size_checks_passed':len(checks),
       'visual_review':{'manuscript_pages':14,'summary_pages':52,
           'all_pages_inspected':True,'render_route':'Hidden Microsoft Word PDF export then PDFium PNG rasterisation',
           'reason_for_fallback':'Packaged renderer reported missing LibreOffice executable',
           'result':'No clipping, overlapping text, broken table borders or unresolved source markers observed'},
       'documents':docs,'independent_human_appraisals':0,
       'exhaustive_systematic_review_complete':False,'journal_submitted':False,
       'drive_project_url':'https://drive.google.com/drive/folders/15-C-RMWbqBHB702RVqLmla5f6XvB_ZPz'}
(REC/'Final_Validation.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')

evidence={p for p in REC.glob('*.md')} | {p for p in REC.glob('*.json')}
for r in register:
    sid=r['source_id']
    for suffix in ('.json','_pages.json'):
        p=REC/'Assessment_Cards'/(sid+suffix)
        if p.is_file():evidence.add(p)
    p=REC/'Crossref_Metadata'/(sid+'.json')
    if p.is_file():evidence.add(p)
for name in ('search_manifest.json','all_records.json','query_01.json','query_02.json','query_03.json',
             'query_04.json','query_05.json','query_06.json','query_07.json','query_08.json',
             'query_09.json','query_10.json','query_11.json','retrieval_register.json'):
    p=REC/'Supplementary_Search'/name
    if p.is_file():evidence.add(p)
proof=REC/'NotebookLM_First_Batch.jpg'
if proof.is_file():evidence.add(proof)
extra=ROOT/'review_2022_2026/search_v3'

def add_evidence(z,base=''):
    for p in sorted(evidence):z.write(p,base+str(p.relative_to(PROJECT)).replace('\\','/'))
    for name in ('search_manifest.json','all_retrieved_candidates.json','assisted_screening.json','FOCUSED_SOURCE_ASSESSMENT.md'):
        p=extra/name
        if p.is_file():z.write(p,base+'Review_Records/Original_Search/'+name)

evidence_zip=PROJECT/'20106UG03_Evidence_Records.zip'
with zipfile.ZipFile(evidence_zip,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:add_evidence(z)
package=OUTPUT/'20106UG03_Review_Package_2026-10-31.zip'
root_files=[p for p in PROJECT.iterdir() if p.is_file() and p.suffix in ('.docx','.md','.csv')]
with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
    for p in root_files:z.write(p,'20106UG03/'+p.name)
    for r in register:
        p=PROJECT/r['file'];z.write(p,'20106UG03/'+str(p.relative_to(PROJECT)).replace('\\','/'))
    add_evidence(z,'20106UG03/')
with zipfile.ZipFile(package) as z:
    assert z.testzip() is None
    assert len([x for x in z.namelist() if x.endswith('.pdf')])==101
    assert len([x for x in z.namelist() if x.endswith('.docx')])==2
print(json.dumps({'package':str(package),'package_bytes':package.stat().st_size,
                  'evidence_zip_bytes':evidence_zip.stat().st_size,'validation':'passed'},indent=2))
