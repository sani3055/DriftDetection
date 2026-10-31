"""Revise the two-author manuscript using saved, exploratory research evidence."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'outputs/SHAP_CloudTrail_Journal_Paper_v4_Two_Authors.docx'
OUTPUT = ROOT/'outputs/CloudTrail_Journal_Manuscript_Guide_Review_Final.docx'
if OUTPUT.exists(): raise FileExistsError('Preserve existing guide-review manuscript')
doc=Document(SOURCE)
static=pd.read_csv(ROOT/'results/journal_v6_static_baselines/metrics.csv')
synthetic=pd.read_csv(ROOT/'results/journal_v5_labelled_generic/summary.csv')
safety=pd.read_csv(ROOT/'results/journal_v5_safety_analysis/descriptive_summary.csv')
assert len(static)==24 and len(synthetic)==20 and len(safety)==200
assert set(static[static.session=='pooled'].events)=={11531}

def find(text):
    matches=[p for p in doc.paragraphs if p.text.startswith(text)]
    if len(matches)!=1: raise ValueError((text,len(matches)))
    return matches[0]

def insert_before(anchor,text,style=None):
    el=OxmlElement('w:p');anchor._p.addprevious(el);p=Paragraph(el,anchor._parent)
    if style: p.style=style
    p.add_run(text)
    return p

def replace_paragraph(prefix,old,new):
    p=find(prefix)
    assert old in p.text and len(p.runs)==1, (prefix, p.text[:100], len(p.runs))
    p.runs[0].text=p.runs[0].text.replace(old,new)

def add_table_before(anchor, headers, rows):
    table=doc.add_table(rows=1, cols=len(headers))
    template=doc.tables[0]
    table.style=template.style
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        edge=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','4'),('space','0'),('color','D9D9D9')]:
            edge.set(qn('w:'+key),value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for i,text in enumerate(headers): table.rows[0].cells[i].text=str(text)
    for row in rows:
        cells=table.add_row().cells
        for i,text in enumerate(row): cells[i].text=str(text)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                for run in paragraph.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(9)
    for cell in table.rows[0].cells:
        tcpr=cell._tc.get_or_add_tcPr(); shading=OxmlElement('w:shd')
        shading.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fill','E6E9ED');tcpr.append(shading)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:run.bold=True
    anchor._p.addprevious(table._tbl)

# Define the abstract abbreviation and use a six-keyword journal-compatible list.
replace_paragraph('Cloud audit anomaly detectors', 'SHAP attribution drift',
                  'SHapley Additive exPlanations (SHAP) attribution drift')
keywords=find('Keywords ')
assert len(keywords.runs)==2
keywords.runs[1].text=keywords.runs[1].text.replace('SHAP; temporal evaluation;', 'SHAP;')

# Add related work before the next section, then renumber all citations by first use.
related_end=find('3 Data and study design')
insert_before(related_end, 'Xu and colleagues study drift detection, SHAP interpretation and adaptive retraining in IoT anomaly detection [21]. Pacheco and Pedroso evaluate an interpretable intrusion-detection lifecycle across 5G planes, including temporal drift and cross-source transfer [22]. These studies do not validate the present CloudTrail policies; they show why the contribution must be a specific, evidence-based comparison rather than the mere combination of explanations and adaptation.')

methods_end=find('5 Results')
insert_before(methods_end, 'A later exploratory static-detector comparison uses the same narrow chronological fitting and held-out records for event-ordinal and combined-ordinal inputs. Isolation Forest, one-class SVM and novelty-mode Local Outlier Factor are fitted after scaling using historical training records only. Each detector uses the higher 95th percentile of its scores on earlier benign calibration records; this matches a calibration rule, not necessarily a realized alert rate. Fixed default hyperparameters are not tuned on the test sessions. This is a static detector comparison, not a matched comparison of adaptation policies.')
insert_before(methods_end, 'A separate candidate-screening amendment uses the frozen detector and earlier benign calibration scores to select candidates no higher than the 25th or 50th percentile, subject to the existing predicted-normal screen. Candidate pool labels construct controlled test conditions and audit outcomes only; the screen and historical guard never receive those labels. This experiment was specified after earlier results on the same deployment were inspected and is therefore exploratory.')

discussion=find('6 Discussion')
insert_before(discussion,'Static detector comparison','Heading 2')
insert_before(discussion,'Table 11 Static detectors under matched historical calibration')
rows=[]
for rep in ('event','combined'):
    for detector,label in [('isolation_forest','Isolation Forest'),('one_class_svm','One-class SVM'),('local_outlier_factor','Local Outlier Factor')]:
        q=static[(static.representation==rep)&(static.detector==detector)&(static.session=='pooled')]
        assert len(q)==1
        row=q.iloc[0]; rows.append([rep.title()+' ordinal',label,f'{row.recall:.3f}',f'{row.false_positive_rate:.3f}',f'{row.f1:.3f}'])
add_table_before(discussion,['Inputs','Static detector','Recall','Benign FPR','F1'],rows)
insert_before(discussion,'The six settings reuse the same 11,531 labelled test events. Calibration uses earlier benign scores and a strict greater-than threshold. Event-ordinal one-class SVM has higher recall than event-ordinal Isolation Forest in this particular setting, but also a much higher benign false-positive rate. Results depend on the representation, threshold and untuned defaults; the comparison does not establish a generally superior detector. Numeric scaling of ordinal categories imposes arbitrary distances, and LOF novelty scoring may be sensitive to repeated feature vectors.')

insert_before(discussion,'Exploratory candidate screening','Heading 2')
point=safety[(safety.configuration=='event_ordinal_narrow_training_offset_s42')&(safety.policy=='historical_q25')&(safety.pool_size==2000)&(safety.nominal_fraction==.2)]
assert len(point)==1
a=point.iloc[0]
insert_before(discussion, f'The 50th-percentile historical-normality filter selected the same candidates as the ordinary frozen screen in all 150 controlled conditions. The stricter 25th-percentile screen changed candidate selection in three of five representations. In event-ordinal pools of 2,000 with a nominal 20 percent attack fraction, it selected an average of {a.selected_candidates:.1f} records including {a.selected_attacks:.1f} attack-labelled records, versus {a.old_screened_candidates:.1f} and {a.old_screened_attacks:.1f}, respectively, under the ordinary screen. These are means over three dependent seeds from one deployment; the filter also excludes benign records. Its later recall and false-positive-rate differences in these conditions were {a.delta_recall:+.3f} and {a.delta_fpr:+.3f}. A lower attack count does not establish a safe or beneficial update.')

insert_before(discussion,'Independent labelled synthetic transfer','Heading 2')
insert_before(discussion,'An independently authored cloud-neutral synthetic sample provides 801 events, including 14 source-labelled anomalies [23]. All twenty seed-42 source models are frozen and applied without target tuning or updates. This is not native CloudTrail: the fixed adapter carries generic action and identity values into corresponding input fields, lacks the original AWS-specific semantics and leaves several inputs missing; missing readOnly is treated as false by the existing history code. The synthetic timestamps predate source training, so this is retrospective cross-domain replay, not a future deployment.')
insert_before(discussion,'Table 12 Frozen-model outcomes on cloud-neutral synthetic events')
rows=[]
for inp,label in [('event_ordinal','Event ordinal'),('event_onehot','Event one hot'),('behaviour_none','Behaviour'),('combined_ordinal','Combined ordinal'),('combined_onehot','Combined one hot')]:
    q=synthetic[synthetic.configuration==inp+'_narrow_training_offset_s42'];assert len(q)==1
    row=q.iloc[0];rows.append([label,f'{int(row.tp)} / 14',f'{int(row.fp)} / 787',f'{row.recall:.3f}',f'{row.false_positive_rate:.3f}'])
add_table_before(discussion,['Inputs','Anomalies alerted','Normal alerts','Recall','Normal alert rate'],rows)
insert_before(discussion,'All-alert models achieve recall one while flagging every synthetic normal event; the event one-hot model flags nothing and has undefined precision. With only fourteen synthetic anomalies and one generator, these results are descriptive and cannot establish AWS attack-detection accuracy, real-world false-positive rate or incident-level generalization. They do not evaluate adaptation or candidate-screening policies.')

replace_paragraph('The comparison uses one anomaly-model family.', 'The comparison uses one anomaly-model family.', 'The adaptation comparison uses one anomaly-model family; the additional detector comparison is static and exploratory.')
replace_paragraph('The value of the current study', 'The value of the current study lies in its controlled comparisons and failure analysis.', 'The value of the current study lies in its controlled comparisons, matched static baselines and failure analysis; the added synthetic transfer experiment does not establish deployment performance.')
replace_paragraph('The project repository is', 'The versioned research evidence is currently retained locally and requires an archived release for independent replication.', 'Code and earlier research evidence are available in the project repository; the new static-baseline, screening and synthetic-transfer outputs require a versioned release before independent replication.')
replace_paragraph('The inherited monthly source', 'The complete local automated test suite contains 93 passing tests.', 'The local automated test suite contained 102 passing tests after the synthetic adapter was added.')

# Preserve existing bibliography styles and insert three verified additional references.
references=find('References')
existing={}
start=next(i for i,p in enumerate(doc.paragraphs) if p._p is references._p)
for p in doc.paragraphs[start+1:]:
    match=re.match(r'^\[(\d+)\]',p.text)
    if match:
        existing[int(match.group(1))]=deepcopy(p._p)
        p._p.getparent().remove(p._p)
assert set(existing)==set(range(1,21))
new_refs={
 21:'Lijuan Xu, Ziyu Han, Dawei Zhao, Xin Li, Fuqiang Yu and Chuan Chen. Addressing Concept Drift in IoT Anomaly Detection: Drift Detection, Interpretation, and Adaptation. IEEE Transactions on Sustainable Computing, 2024, 9(6), 913–924. https://doi.org/10.1109/TSUSC.2024.3386667',
 22:'Edson José Pacheco and Carlos Marcelo Pedroso. Operational lifecycle of interpretable multi-plane intrusion detection in 5G networks. Computer Networks, 2026, 289, 112672. https://doi.org/10.1016/j.comnet.2026.112672',
 23:'Kateryna Oblakevych. Cloud Threat Anomaly Lab (CTAL), version 0.1.0. Software repository, 2026. https://github.com/katerina20/cloud-threat-anomaly-lab',
}
seen=[]
pattern=re.compile(r'\[(\d+(?:\s*,\s*\d+)*)\]')
for p in doc.paragraphs:
    for match in pattern.finditer(p.text):
        for val in map(int,re.findall(r'\d+',match.group(1))):
            if val not in seen:seen.append(val)
assert set(seen)==set(range(1,24)),seen
mapping={old:i+1 for i,old in enumerate(seen)}
def renumber(match):return '['+', '.join(str(mapping[int(v)]) for v in re.findall(r'\d+',match.group(1)))+']'
for p in doc.paragraphs:
    if pattern.search(p.text):
        assert len(p.runs)==1, ('citation spans multiple runs',p.text[:100])
        p.runs[0].text=pattern.sub(renumber,p.runs[0].text)
last=references._p
for old in seen:
    if old in existing:
        element=existing[old]
        text_nodes=element.xpath('.//w:t')
        assert text_nodes and text_nodes[0].text.startswith(f'[{old}]')
        text_nodes[0].text=re.sub(r'^\[\d+\]',f'[{mapping[old]}]',text_nodes[0].text,count=1)
    else:
        element=OxmlElement('w:p')
        p=Paragraph(element,references._parent)
        p.style=doc.styles['Normal'];p.paragraph_format.left_indent=Pt(15);p.paragraph_format.first_line_indent=Pt(-15)
        run=p.add_run(f'[{mapping[old]}] '+new_refs[old]);run.font.name='Times New Roman';run.font.size=Pt(9.5)
    last.addnext(element);last=element

doc.core_properties.author='Sanidhya Bhandari; Badrinath N'
doc.save(OUTPUT)
out=ROOT/'tmp/journal_guide_review';out.mkdir(parents=True,exist_ok=True)
audit={'output':str(OUTPUT),'sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
       'citations_in_first_use_order':seen,'bibliography_entries':len(seen),
       'source_docx_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
       'evidence_sha256':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [
           ROOT/'results/journal_v6_static_baselines/metrics.csv',ROOT/'results/journal_v5_labelled_generic/summary.csv',
           ROOT/'results/journal_v5_safety_analysis/descriptive_summary.csv']}}
(out/'build_audit.json').write_text(json.dumps(audit,indent=2)+'\n',encoding='utf-8')
print(json.dumps(audit,indent=2))
