"""Build an evidence-linked journal manuscript without changing experiment outputs.

Run figures with the research environment, then run this builder with the bundled
document Python. The intermediate DOCX is rendered and inspected before delivery.
"""
from __future__ import annotations

from copy import deepcopy
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
TMP=ROOT/'tmp/journal_paper'
OUT=ROOT/'outputs/SHAP_CloudTrail_Journal_Paper.docx'
SOURCE=ROOT/'docs/journal_paper_source.md'
RESOLVED=ROOT/'docs/journal_paper.md'
AUDIT=TMP
REFERENCES=json.loads((ROOT/'docs/journal_references.json').read_text(encoding='utf-8'))
FONT='Times New Roman'
WIDTH=6.5
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}


def read_csv(relative):
    with (ROOT/relative).open(encoding='utf-8',newline='') as handle:
        return list(csv.DictReader(handle))


def read_json(relative):
    return json.loads((ROOT/relative).read_text(encoding='utf-8'))


def file_hash(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def f4(value):
    return f'{float(value):.4f}'


def f6(value):
    return f'{float(value):.6f}'.replace('-','−')


def font(run,size=None,bold=None):
    run.font.name=FONT
    run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),FONT)
    run._element.rPr.rFonts.set(qn('w:hAnsi'),FONT)
    if size is not None:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    return run


doc=Document()
section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.8)
section.left_margin=section.right_margin=Inches(1)
section.footer_distance=Inches(.35)
for name,size in [('Normal',11),('Title',17),('Heading 1',13),('Heading 2',11.5),('Caption',9.5)]:
    style=doc.styles[name]
    style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'),FONT)
    style._element.rPr.rFonts.set(qn('w:hAnsi'),FONT)
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.12
    if name.startswith('Heading'):
        style.font.bold=True
        style.paragraph_format.space_before=Pt(11)
        style.paragraph_format.keep_with_next=True
    if name=='Title':style.font.bold=True
doc.styles['Normal'].paragraph_format.widow_control=True
doc.styles['Caption'].font.bold=False
doc.styles['Caption'].font.italic=False
# The bundled Word template can contain a blue paragraph border on Title.
# Remove paragraph borders from this new manuscript's styles, not table borders.
for border in doc.styles._element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.core_properties.title='SHAP Guided Drift Diagnosis for Adaptive Anomaly Detection in AWS CloudTrail'
doc.core_properties.author='Sanidhya Bhandari'
doc.core_properties.subject='Research article on attribution feature coupling and drift-guided updates'

footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);field.append(r)
footer._p.append(field)

used=[]
resolved_blocks=[]
table_audit={}
equation_count=0


def citations(text):
    def replace(match):
        keys=re.findall(r'@([a-z]+)',match.group(0))
        numbers=[]
        for key in keys:
            if key not in REFERENCES:raise KeyError(key)
            if key not in used:used.append(key)
            numbers.append(str(used.index(key)+1))
        return '['+', '.join(numbers)+']'
    return re.sub(r'\[@[^\]]+\]',replace,text)


def paragraph(text,style=None,align=None):
    text=citations(text)
    p=doc.add_paragraph(style=style)
    if align is not None:p.alignment=align
    elif not style:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for i,part in enumerate(re.split(r'\*\*([^*]+)\*\*',text)):
        if part:font(p.add_run(part),bold=True if i%2 else None)
    resolved_blocks.append(text)
    return p


def table(key,title,headers,rows,widths,note=''):
    number=len(table_audit)+1
    caption=paragraph(f'Table {number} {title}','Caption')
    caption.paragraph_format.keep_with_next=True
    caption.paragraph_format.space_before=Pt(7)
    for r in caption.runs:r.bold=True
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        node=OxmlElement(f'w:{edge}');node.set(qn('w:val'),'single');node.set(qn('w:sz'),'4');node.set(qn('w:color'),'D9D9D9');borders.append(node)
    props.append(borders)
    grid=t._tbl.tblGrid
    for node,w in zip(grid,widths):node.set(qn('w:w'),str(round(w*1440)))
    for i,rowvalues in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value,w) in enumerate(zip(cells,rowvalues,widths)):
            cell.width=Inches(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBF0' if i==0 else 'FFFFFF');pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                m=OxmlElement(f'w:{edge}');m.set(qn('w:w'),'75' if edge in ['left','right'] else '65');m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            p=cell.paragraphs[0]
            # Keep compact comparison tables together instead of orphaning their
            # final row on another page. Longer tables may continue with headers.
            if len(rows)<=8:
                p.paragraph_format.keep_with_next=i<len(rows)
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or key in ['protocol','scenarios','diagnosis'] else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(str(value)),9.5,i==0)
    if note:
        p=paragraph(note)
        p.paragraph_format.space_before=Pt(5)
        for r in p.runs:r.font.size=Pt(9)
    else:
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
        font(p.add_run(''),4)
    table_audit[key]={'number':number,'title':title,'headers':headers,'rows':rows}
    resolved_blocks.append('\n'.join(['| '+' | '.join(map(str,headers))+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(map(str,row))+' |' for row in rows]))


def figure(key):
    settings={
      'pipeline':('pipeline.png',6.45,'Figure 1 Offline monitoring and update logic. Feature distributions and model attributions provide separate evidence. Labels do not enter the runtime gates; candidate normality does not guarantee an uncontaminated refit.'),
      'signals':('controlled_signals.png',6.1,'Figure 2 Held-out controlled scenario signals. Marker shapes and colours identify intended scenario families, not verified incidents. Dashed lines show magnitude gates selected on controlled validation examples. Repeated markers can overlap; alignment is an additional gate not displayed on these axes.'),
      'native':('native_prequential.png',6.1,'Figure 3 Native chronological evaluation at 128 events per window. Rates pool confusion counts over future events. Coupled policies coincide with static because their gates request no updates. These are descriptive results from one controlled deployment, not independent production incidents.'),
    }
    filename,width,caption=settings[key]
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
    run=p.add_run();run.add_picture(str(TMP/'assets'/filename),width=Inches(width))
    inline=run._r.xpath('.//wp:docPr')[0];inline.set('descr',caption)
    p=paragraph(caption,'Caption');p.paragraph_format.keep_with_next=False
    p.paragraph_format.space_after=Pt(7)


# Native, editable Office Math. Each argument can be a string or an OMML node.
def mathrun(text):
    node=OxmlElement('m:r')
    rpr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');rpr.append(fonts)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'22');rpr.append(size);node.append(rpr)
    t=OxmlElement('m:t');t.text=text;t.set(qn('xml:space'),'preserve');node.append(t)
    return node


def append(parent,items):
    for item in items if isinstance(items,list) else [items]:parent.append(mathrun(item) if isinstance(item,str) else deepcopy(item))


def sub(base,index):
    n=OxmlElement('m:sSub');e=OxmlElement('m:e');s=OxmlElement('m:sub');append(e,base);append(s,index);n.append(e);n.append(s);return n


def sup(base,index):
    n=OxmlElement('m:sSup');e=OxmlElement('m:e');s=OxmlElement('m:sup');append(e,base);append(s,index);n.append(e);n.append(s);return n


def frac(numerator,denominator):
    n=OxmlElement('m:f');a=OxmlElement('m:num');b=OxmlElement('m:den');append(a,numerator);append(b,denominator);n.append(a);n.append(b);return n


def summation(lower,upper,expression):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr');char=OxmlElement('m:chr');char.set(qn('m:val'),'∑');pr.append(char)
    loc=OxmlElement('m:limLoc');loc.set(qn('m:val'),'subSup');pr.append(loc);n.append(pr)
    for tag,items in [('sub',lower),('sup',upper),('e',expression)]:
        e=OxmlElement(f'm:{tag}');append(e,items);n.append(e)
    return n


def delimiter(expression,character='|'):
    n=OxmlElement('m:d');pr=OxmlElement('m:dPr')
    for tag in ['begChr','endChr']:
        c=OxmlElement(f'm:{tag}');c.set(qn('m:val'),character);pr.append(c)
    n.append(pr);e=OxmlElement('m:e');append(e,expression);n.append(e)
    return n


def equation(key):
    global equation_count
    equation_count+=1
    if key=='psi':
        parts=[sub('d','j'),' = ',summation(['b ∈ ',sub('B','j')],'',[
            '(',sub('q','jb'),' − ',sub('p','jb'),') ln ',frac(sub('q','jb'),sub('p','jb'))])]
    elif key=='feature':
        parts=[sub('F','t'),' = ',frac('1','m'),summation('j = 1','m',sub('d','j,t')),' + ',frac(sub('z','t'),'10')]
    elif key=='wasserstein':
        parts=[sub('w','j,t'),' = ',frac([sub('W','1'),'(',sup(sub('Φ','j'),'R'),', ',sup(sub('Φ','j'),'t'),')'],[sup(sub('σ','j'),'R'),' + ',sup('10','−9')])]
    elif key=='shap':
        parts=[sub('D','t'),' = ',frac('1','p'),summation('j = 1','p',sub('w','j,t')),' + (1 − ',sub('J','t'),') + (1 − ',frac([sub('K','t'),' + 1'],'2'),')']
    elif key=='coupling':
        parts=[sub('a','j,t'),' = ',delimiter([sub('E','t'),delimiter(sub('φ','j')),' − ',sub('E','R'),delimiter(sub('φ','j'))]),',     ',sub('C','t'),' = ',frac([sup(sub('d','t'),'T'),sub('a','t')],[sub(delimiter(sub('d','t'),'‖'),'2'),' ',sub(delimiter(sub('a','t'),'‖'),'2')])]
    elif key=='cost':
        parts=['L = ',sub('c','m'),' ',sub('1','sec'),'(1 − r) + ',sub('c','f'),' ',sub('1','benign'),' FPR + ',sub('c','u'),' u']
    else:raise KeyError(key)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(7)
    math=OxmlElement('m:oMath');append(math,parts);p._p.append(math)
    font(p.add_run(f'    ({equation_count})'),10)
    resolved_blocks.append(f'[Editable equation {equation_count} {key}]')


def build_native_table(key):
    if key=='native_policies':
        rows=[r for r in read_csv('results/journal_v2/w128_q0.95_s42/policy_summary.csv') if r['session']=='pooled']
        table(key,'Prequential detection on native future events',['Policy','Recall','F1','FPR','Requests','Refits'],
              [[r['policy'].replace('_',' ')]+[f4(r[c]) for c in ['recall','f1','false_positive_rate']]+[r['update_requests'],r['completed_refits']] for r in rows],
              [2.0,.9,.9,.9,.9,.9], 'All seven policies evaluate the same 11,531 events, including 3,126 attack-labelled and 8,405 benign events. Scoring precedes updating. A request can fail when fewer than ten baseline-predicted normal candidates exist. Terminal-window requests cannot improve a subsequent window in that session.')
    elif key=='native_sensitivity':
        paths=sorted((ROOT/'results/journal_v2').glob('*/manifest.json'))
        body=[]
        for path in paths:
            manifest=json.loads(path.read_text(encoding='utf-8'))
            params=manifest['parameters']
            policies={r['policy']:r for r in read_csv(str((path.parent/'policy_summary.csv').relative_to(ROOT))) if r['session']=='pooled'}
            body.append([f"{params['window_size']} / {params['quantile']} / {params['seed']}",f4(policies['static']['f1']),f4(policies['dual_signal']['f1']),f4(policies['static']['false_positive_rate']),f4(policies['dual_signal']['false_positive_rate']),policies['coupled_signal']['update_requests']])
        if len(body)!=7:raise ValueError('All seven locked sensitivity settings must finish before manuscript export')
        table(key,'Full native protocol sensitivity grid',['Window / q / seed','Static F1','Dual F1','Static FPR','Dual FPR','Coupled requests'],body,[2.0,.9,.9,.9,.9,.9],
              'q is the benign calibration quantile. The 128 / 0.95 / 42 setting remains primary. Other settings were declared as sensitivity analyses after the primary run. All reuse the same source deployment and cannot be treated as independent replications.')
    else:raise KeyError(key)


def build_table(key):
    if key.startswith('native_'):
        return build_native_table(key)
    if key=='protocol':
        table(key,'Data tracks and distinct evaluation units',['Track','Data and partition','Permitted interpretation'],[
            ['Primary corpus','1,939,207 rows; 45 months; no attack labels','Observed distribution and attribution change'],
            ['Chronological split','27 training / 9 validation / 9 test months','Time-respecting real-data observation'],
            ['Controlled tests','35 validation / 70 test scenario–seed units','Intervention diagnosis and current-window scoring'],
            ['External detector','16,495 benign training / 24,087 labelled test events','Event-level detector evaluation'],
            ['External checkpoint','10 security groups + benign target + benign reference','Retrospective group-level diagnosis; not future recovery'],
        ],[1.25,2.4,2.85])
    elif key=='scenarios':
        table(key,'Controlled intervention families',['Family','Field replacements','Selected fractions'],[
          ['Stable','No replacement','0'],
          ['Benign region','awsRegion = ap-southeast-1','25%'],
          ['Benign identity','userIdentitytype = AWSService','25%'],
          ['Benign service','eventSource = lambda.amazonaws.com','25%'],
          ['Security privilege','eventName = AttachUserPolicy; identity type = Root','25%; 12.5%'],
          ['Security defence','eventName = DeleteTrail; service = cloudtrail.amazonaws.com','25%; 12.5%'],
          ['Security credentials','eventName = AssumeRole; region = ap-southeast-1','25%; 12.5%'],
        ],[1.25,4.05,1.2], 'Fractions refer to separate selected-index sets. They need not be disjoint or represent distinct changed values. Category replacements use the training map, with unseen values mapped to −1.')
    elif key=='diagnosis':
        table(key,'Applied coupling diagnosis rule',['Feature gate','SHAP gate','Alignment gate','Output and action'],[
            ['Low','Low','Any','Stable; no request'],['High','Low','Any','Benign drift; no request'],
            ['High','High','High','Security-relevant drift; request update'],
            ['High','High','Low','Review required; no request'],
            ['Low','High','Low / High','Instability / review required; no request'],
        ],[.85,.85,1.0,3.8], 'High means at least the applicable threshold. Diagnosis names describe model-derived states; they are not confirmed incident labels.')
    elif key=='policies':
        values={r['policy']:r for r in read_csv('data/benchmark/adaptation_summary.csv')}
        names=[('static','Static'),('periodic','Periodic'),('feature_only','Feature only'),('shap_only','SHAP only'),('dual_signal','Dual'),('coupled_signal','Coupled'),('conservative_coupled_signal','Conservative coupled')]
        rows=[[name]+[f4(values[k][col]) for col in ['mean_f1','mean_recall','mean_false_positive_rate','trigger_rate','unnecessary_adaptation_rate']] for k,name in names]
        table(key,'Held-out current-window policy results',['Policy','Mean F1','Mean recall','Mean FPR','Update rate','Unnecessary rate'],rows,[1.6,.8,.85,.85,.9,1.5], 'Each mean covers all 70 units, including four scenario families with no positives. Both rates divide by all runs. These are controlled selected-row targets, not real attack labels.')
    elif key=='paired':
        rows=read_csv('results/paired_inference.csv')
        values={r['metric']:r for r in rows}
        body=[]
        for key_,name in [('false_positive_rate','False-positive rate'),('f1','Mean event F1')]:
            r=values[key_]
            body.append([name,f6(r['mean_difference']),f"[{f6(r['ci_low'])}, {f6(r['ci_high'])}]",f"{float(r['q_value_bh']):.7f}" if key_=='false_positive_rate' else f6(r['q_value_bh'])])
        table(key,'Paired dual minus static differences',['Measure','Mean difference','Bootstrap 95% interval','BH adjusted q'],body,[1.3,1.25,2.65,1.3], 'Paired units share scenario and seed. Negative FPR is favourable. The FPR Wilcoxon p-value is 0.0002896194. The F1 interval does not establish equivalence or improvement.')
    elif key=='detectors':
        rows=read_csv('data/robustness/detector_ablation_summary.csv')
        names={'isolation_forest_50':'Isolation Forest 50 trees','isolation_forest_100':'Isolation Forest 100 trees','isolation_forest_200':'Isolation Forest 200 trees','local_outlier_factor':'Local Outlier Factor','one_class_svm':'One-Class SVM'}
        table(key,'Separate controlled detector robustness results',['Detector','Security recall','Security F1','Benign FPR'],[[names[r['detector']]]+[f4(r[c]) for c in ['security_recall','security_f1','benign_false_positive_rate']] for r in rows],[2.6,1.3,1.3,1.3], 'Detectors fit a 200-row prefix in this experiment. Security metrics average the three security families; benign FPR averages stable and benign families. This is not the policy-comparison fitting protocol.')
    elif key=='external_detectors':
        names={'isolation_forest':'Isolation Forest','local_outlier_factor':'Local Outlier Factor','one_class_svm':'One-Class SVM'}
        rows=read_csv('data/external_validation/external_detector_metrics.csv')
        table(key,'Independent event-level detector evaluation',['Detector','Precision','Recall','F1','FPR'],[[names[r['detector']]]+[f4(r[c]) for c in ['precision','recall','f1','false_positive_rate']] for r in rows],[1.8,1.175,1.175,1.175,1.175], 'Training uses the separate benign archive. Attack and benign labels in the mixed test archive are used for event-level evaluation.')
    elif key=='external_windows':
        observations=read_csv('data/external_validation/external_temporal_diagnosis.csv')
        metrics=read_json('data/external_validation/external_temporal_diagnosis_metrics.json')
        selections=[('dual_signal','dual_signal_prediction','Feature / SHAP / dual'),('coupled_signal','coupled_signal_prediction','Source cosine coupling'),('conservative_coupled_signal','conservative_coupled_signal_prediction','Source conservative coupling'),('unsupervised_coupled_signal','unsupervised_coupled_signal_prediction','Target batch recalibration')]
        body=[]
        for policy,col,name in selections:
            pairs=[(int(r['ground_truth_binary']),int(r[col])) for r in observations]
            counts=[sum(y==1 and p==1 for y,p in pairs),sum(y==0 and p==1 for y,p in pairs),sum(y==0 and p==0 for y,p in pairs),sum(y==1 and p==0 for y,p in pairs)]
            body.append([name]+counts+[f4(metrics[policy][m]) for m in ['recall','f1','false_positive_rate']])
        table(key,'Retrospective external group-level checkpoint',['Rule','TP','FP','TN','FN','Recall','F1','FPR'],body,[2.45,.45,.45,.45,.45,.75,.75,.75], 'TP, FP, TN and FN are group counts. The 12 units include the benign reference. Target batch recalibration uses the conservative alignment score and the same batch for threshold estimation and evaluation.')
    else:raise KeyError(key)


def add_references():
    for number,key in enumerate(used,1):
        ref=REFERENCES[key]
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Inches(.28);p.paragraph_format.first_line_indent=Inches(-.28)
        p.paragraph_format.space_after=Pt(6);p.paragraph_format.line_spacing=1.0
        font(p.add_run(f"[{number}] {ref['authors']}. {ref['title']}. "),9.5)
        font(p.add_run(ref['venue']),9.5).italic=True
        tail=f", {ref['year']}"
        if ref.get('volume'):tail+=f", {ref['volume']}"+(f"({ref['issue']})" if ref.get('issue') else '')
        if ref.get('pages'):tail+=f", {ref['pages']}"
        tail+='.'
        if ref.get('accessed'):tail+=f" Accessed {ref['accessed']}."
        font(p.add_run(tail+' '+ref['url']),9.5)
        resolved_blocks.append(p.text)


def validate_evidence():
    manifest=read_json('data_manifest/real_cloudtrail_manifest.json')
    assert manifest['row_count']==1939207
    assert manifest['sha256']=='b3aa4fcf28747f4f84457d2d3b76f4c27901bd2e6289bf6d6288cb9587d20589'
    assert read_json('models/isolation_forest_event_baseline.json')['training_rows']==251613
    policies={r['policy']:r for r in read_csv('data/benchmark/adaptation_summary.csv')}
    assert f4(policies['dual_signal']['mean_false_positive_rate'])=='0.3397'
    assert f4(policies['static']['mean_false_positive_rate'])=='0.4282'
    assert f4(read_json('data/benchmark/diagnosis_thresholds.json')['validation_macro_f1'])=='0.7306'
    assert f6(read_json('data/external_validation/external_temporal_diagnosis_manifest.json')['unsupervised_transfer_threshold'])=='0.499080'
    diag=[r for r in read_csv('data/benchmark/diagnosis_results.csv') if r['split']=='test']
    assert len(diag)==70
    assert all(r['diagnosis']==r['coupled_diagnosis']==r['conservative_coupled_diagnosis'] for r in diag)
    assert sum(r['scenario']=='security_privilege' and r['coupled_diagnosis']=='security_relevant_drift' for r in diag)==0
    assert sum(r['scenario']=='security_defense' and r['coupled_diagnosis']=='security_relevant_drift' for r in diag)==7
    adaptation=read_csv('data/benchmark/adaptation_results.csv')
    for policy,requests,unnecessary in [('dual_signal',17,0),('feature_only',53,23),('periodic',24,14)]:
        rows=[r for r in adaptation if r['policy']==policy]
        assert sum(int(r['triggered']) for r in rows)==requests
        assert sum(int(r['unnecessary_adaptation']) for r in rows)==unnecessary


def main():
    global SOURCE,OUT,RESOLVED,AUDIT
    parser=argparse.ArgumentParser()
    parser.add_argument('--source',type=Path,default=SOURCE)
    parser.add_argument('--output',type=Path,default=OUT)
    parser.add_argument('--resolved',type=Path,default=RESOLVED)
    parser.add_argument('--audit-dir',type=Path,default=AUDIT)
    args=parser.parse_args()
    SOURCE,OUT,RESOLVED,AUDIT=args.source,args.output,args.resolved,args.audit_dir
    AUDIT.mkdir(parents=True,exist_ok=True)
    if 'v2' in SOURCE.stem:
        REFERENCES['lara']={'authors':'Antonio Lara-Gutierrez, Carmen Fernandez-Gago and Jose A. Onieva',
            'title':'A Framework for Drift Detection and Adaptation in AI-driven Anomaly and Threat Detection Systems',
            'venue':'International Journal of Information Security','year':'2025','volume':'24','issue':'5',
            'pages':'199','url':'https://doi.org/10.1007/s10207-025-01118-9','kind':'journal'}
    TMP.mkdir(parents=True,exist_ok=True)
    watched=[ROOT/p for p in ['data/benchmark/adaptation_results.csv','data/benchmark/diagnosis_results.csv','data/benchmark/diagnosis_thresholds.json','results/paired_inference.csv','data/external_validation/external_temporal_diagnosis_metrics.json']]
    before={str(p.relative_to(ROOT)):file_hash(p) for p in watched}
    validate_evidence()
    source=SOURCE.read_text(encoding='utf-8')
    for block in source.strip().split('\n\n'):
        block=block.strip()
        if not block:continue
        if block.startswith('# '):
            p=paragraph(block[2:],'Title',WD_ALIGN_PARAGRAPH.CENTER)
            p.paragraph_format.space_after=Pt(9)
        elif block.startswith('### '):
            paragraph(re.sub(r'^\d+\.\d+\s+','',block[4:]),'Heading 2')
        elif block.startswith('## '):paragraph(block[3:],'Heading 1')
        elif block.startswith('[[TABLE:'):build_table(block[8:-2])
        elif block.startswith('[[FIG:'):figure(block[6:-2])
        elif block.startswith('[[EQ:'):equation(block[5:-2])
        elif block=='[[REFERENCES]]':add_references()
        elif block=='Sanidhya Bhandari' or block.startswith('Department of IoT,'):
            p=paragraph(block,align=WD_ALIGN_PARAGRAPH.CENTER)
            if block=='Sanidhya Bhandari':
                for r in p.runs:r.bold=True
            else:
                for r in p.runs:r.font.size=Pt(10)
        else:paragraph(block.replace('\n',' '))
    expected_keys=set(re.findall(r'@([a-z]+)',source))
    assert set(used)==expected_keys,(used,expected_keys)
    assert len(table_audit)==source.count('[[TABLE:')
    assert equation_count==source.count('[[EQ:')
    OUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUT)
    final_text='\n\n'.join(resolved_blocks)
    assert not re.search(r'\[\[|\[@',final_text)
    RESOLVED.write_text(final_text+'\n',encoding='utf-8')
    assert before=={str(p.relative_to(ROOT)):file_hash(p) for p in watched}
    audit={'output':str(OUT),'source_words':len(source.split()),'tables':table_audit,'editable_equations':equation_count,
           'references_in_citation_order':used,'preserved_research_artifacts':before,'output_sha256':file_hash(OUT)}
    (AUDIT/'build_audit.json').write_text(json.dumps(audit,indent=2),encoding='utf-8')
    print(json.dumps({k:audit[k] for k in ['output','source_words','editable_equations','references_in_citation_order','output_sha256']},indent=2))


if __name__=='__main__':main()
