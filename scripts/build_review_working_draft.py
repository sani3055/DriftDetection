"""Create a transparent review draft, not a falsely completed 100-study review."""
from collections import Counter
from datetime import date
import html
import json
from pathlib import Path
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'review_2022_2026/submission_audit'
OUT = ROOT / 'outputs/CloudTrail_PRISMA_Review_Working_Draft.docx'


def main():
    records = json.loads((AUDIT / 'verified_candidates.json').read_text(encoding='utf-8'))
    manifest = json.loads((AUDIT / 'audit_manifest.json').read_text(encoding='utf-8'))
    counts = manifest['counts']
    years = Counter(row['verified_first_publication_date'][:4] for row in records)
    abstract_count = sum(bool(row.get('verified_abstract')) for row in records)
    journals = len({row['verified_journal'] for row in records})
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(.8)
    sec.left_margin = sec.right_margin = Inches(1)
    for name, size in [('Normal', 11), ('Title', 17), ('Subtitle', 11), ('Heading 1', 13), ('Heading 2', 11.5)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.12
        if name.startswith('Heading'):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
    for border in doc.styles._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.core_properties.author = 'Sanidhya Bhandari'
    doc.core_properties.title = 'Explainable Drift Monitoring for Adaptive Cybersecurity Detection'
    blocks = []

    def paragraph(text, style=None):
        doc.add_paragraph(text, style)
        blocks.append(text)

    def heading(text):
        paragraph(text, 'Heading 1')

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        border = OxmlElement('w:tblBorders')
        for side in ['top','bottom','left','right','insideH','insideV']:
            item=OxmlElement('w:'+side)
            for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]: item.set(qn('w:'+key),value)
            border.append(item)
        t._tbl.tblPr.append(border)
        for col,width in zip(t.columns,widths): col.width=Inches(width)
        for i, values in enumerate([headers]+rows):
            cells = t.rows[0].cells if i == 0 else t.add_row().cells
            props = cells[0]._tc.getparent().get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if i==0: props.append(OxmlElement('w:tblHeader'))
            for cell,value,width in zip(cells,values,widths):
                cell.width=Inches(width)
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text=str(value)
                margins=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),'75');el.set(qn('w:type'),'dxa');margins.append(el)
                cell._tc.get_or_add_tcPr().append(margins)
                if i==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EBF0');cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    if len(rows)<=8:
                        p.paragraph_format.keep_with_next=i<len(rows)
                    p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.line_spacing=1.05
                    for run in p.runs:
                        run.font.name='Times New Roman';run.font.size=Pt(10)
                        if i==0:run.bold=True
        doc.add_paragraph()
        blocks.extend([' | '.join(map(str, row)) for row in [headers]+rows])

    paragraph('Explainable Drift Monitoring for Adaptive Cybersecurity Detection', 'Title')
    paragraph('Systematic review manuscript in preparation', 'Subtitle')
    paragraph('Sanidhya Bhandari')
    paragraph('Department of IoT, School of Computer Science and Engineering, Vellore Institute of Technology, Vellore, India')
    heading('Abstract')
    paragraph(f'''Security detectors face changing workloads, but a distribution shift does not by itself identify an attack or justify retraining. This review examines how recent journal research connects anomaly detection, feature attribution and adaptation, with cloud audit logs as the motivating application. The planned synthesis separates direct cloud-audit evidence from adjacent network, IoT and general log research and assesses temporal validity, label use and reproducibility. A metadata search conducted on 21 October 2026 produced 100 initial candidates. Directed primary-source follow-up and DOI verification on 29 October expanded the current register to {len(records)} unique journal-article records with reported publication years from 2022 to 2026. {counts['readable_pdf_files_verified']} candidate records have readable PDF files available locally. These are acquisition counts, not the number of included studies. Full-text eligibility and study-level critical appraisal remain incomplete, so the review does not yet report a final PRISMA flow, an included-study synthesis or comparative performance conclusions. The proposed evidence framework distinguishes drift measurement, explanation change and demonstrated future recovery after adaptation. It provides a basis for assessing whether attribution-guided update claims survive time-respecting evaluation.''')
    paragraph('Keywords: systematic review; cloud audit logs; explainable artificial intelligence; concept drift; anomaly detection; adaptation; PRISMA')
    heading('1 Introduction')
    paragraph('Cloud audit logs record actions taken by users, services and automated workloads. A detector can flag unfamiliar activity when a workload changes, even when the activity is legitimate. Conversely, an attacker can use common administrative operations. These conditions make the interpretation of changing anomaly scores important for cloud monitoring.')
    paragraph('The research problem extends beyond detecting input drift. A system must determine whether a change matters to its decisions, whether an explanation helps interpret the change, and whether updating improves subsequent predictions. These are distinct evidential questions. A study that explains an alert does not necessarily evaluate adaptation, and a study that lowers error on its update batch does not necessarily demonstrate future recovery.')
    paragraph('The review will examine journal research from 2022 to 2026. It uses cloud audit monitoring as its focal application while retaining clearly identified adjacent security methods. This broader scope avoids implying that every retrieved article studies AWS CloudTrail. Direct and adjacent evidence will remain separate throughout extraction and synthesis.')
    heading('2 Review questions')
    paragraph('RQ1 asks which datasets, representations and label assumptions support cloud or adjacent security anomaly detection. RQ2 asks how studies measure input drift and changes in explanations. RQ3 asks whether those signals affect an adaptation policy and whether the study evaluates later, untouched data. RQ4 asks which evidence supports transfer, contamination control and operational reproducibility.')
    paragraph('The intended contribution is a structured account of the evidence supporting an update decision. It will distinguish a change in model attribution from a verified security incident or a measured benefit of adaptation. The categories below are proposed extraction dimensions, not findings about all candidate papers.')
    heading('3 Methods')
    paragraph('The reporting framework follows PRISMA 2020 [1]. The supplied checklist governs the reporting items, including search sources and dates, selection, extraction, study appraisal, synthesis and access to materials. PRISMA does not prescribe a minimum of 100 studies. That number is the capstone requirement, and final inclusion must still follow relevance and quality checks. The protocol is not externally registered.')
    paragraph('Eligibility covers journal articles first published from 1 January 2022 through 29 October 2026, with a verifiable DOI and a substantive connection to security detection and at least one of explanation, drift, adaptation or cloud/log analysis. Conference papers may support background discussion but will not count toward the required journal corpus. Retractions, nonresearch notices and unrelated applications are excluded after verification. Existing surveys will be distinguished from primary experimental studies rather than mixed as equivalent evidence.')
    paragraph('The initial metadata discovery used OpenAlex with a Crossref fallback. The archived queries cover AWS CloudTrail, cloud audit anomaly detection, explainable intrusion detection, SHAP drift, concept drift and adaptive cybersecurity. A directed follow-up added six relevant journal records identified from primary publisher pages and the research bibliography. That follow-up is a separate discovery route and must appear separately in the eventual flow accounting.')
    paragraph('The initial collector saved only its selected candidate set rather than a complete record-level export of every search hit. Consequently, the historical aggregate hit counts cannot be reconstructed into a fully audited exclusion flow from the current files. A database search update must archive all retrieved records and exact database-specific syntax before a final review is submitted. Scopus, Web of Science, IEEE Xplore and ACM searches have not been claimed as completed.')
    paragraph('The DOI verification step queries Crossref and checks article type, bibliographic identity and first available publication date. It records publisher licence links and retrieves PDFs from exposed links when an explicit Creative Commons licence is present. Previously available files are checked separately. PDF parsing, page counts, hashes and a DOI search on the first two pages help check file identity, but parsing alone does not prove that the file is the correct final article, that it has been critically read, or that redistribution is permitted.')
    paragraph('Screening will use title and abstract first, followed by full-text eligibility with one recorded reason for each exclusion. AI-assisted triage and data extraction must retain their evidence locations. No completed independent human screening round is claimed. PRISMA requires an accurate account of who screened and how automation was used. A second human check would reduce selection error, but an AI pass must not be described as an independent human reviewer. The author remains responsible for the final manuscript.')
    paragraph('For each included paper, extraction will record the dataset and its provenance, input representation, label access, detector, explanation method, drift metric, calibration, adaptation strategy, baseline settings, evaluation chronology, metrics and uncertainty, external validation, code/data access and stated limitations. Missing values will remain not reported rather than being inferred. Claims about false positives or attack recall require the paper to define the corresponding denominator and ground truth.')
    heading('4 Evidence framework')
    table(['Dimension','Question asked of each study'],[
        ['Security target','Does the evaluation use true attack labels, anomaly labels, selected-row masks or no ground truth?'],
        ['Time order','Are training, calibration and evaluation periods separated? Do updates precede the records used to evaluate them?'],
        ['Explanation evidence','What model output does the explanation describe? Is reconstruction or stability checked?'],
        ['Adaptation evidence','Does the method compare updates with a static and a simple update baseline under matched settings?'],
        ['Transfer and access','Does the study test another account or dataset? Can code, data and calibration decisions be traced?'],
    ],[1.5,5.0])
    paragraph('Temporal evidence will be coded descriptively as random/static evaluation, chronological evaluation, test-then-update evaluation with future recovery, or evaluation across independent environments. These are analysis categories, not a validated numerical quality scale. A chronological split can coexist with leakage if future labels influence threshold selection or if the update batch also supplies the recovery score.')
    paragraph('The synthesis will group papers by direct cloud-audit monitoring, explainable security detection, drift/adaptation and general log methods. Results from different datasets will not be ranked by raw accuracy. A narrative synthesis is planned because label definitions, imbalance, model outputs and units of analysis are heterogeneous. No meta-analysis or pooled effect estimate has been performed.')
    paragraph('Critical appraisal will consider dataset and label validity, chronological leakage, baseline comparability, test-informed tuning, same-batch recovery claims, independence of repeated trials and reproducibility. Each judgment needs a full-text location and an explanation. Citation count, a DOI or a journal name alone cannot establish methodological quality. Reporting-bias and certainty judgments remain pending rather than assumed favourable.')
    heading('5 Acquisition results')
    paragraph(f'The current register contains {len(records)} unique DOI records from {journals} journal titles. Crossref confirms their journal-article type and a reported publication year within 2022 to 2026. Publication-date bounds establish the search cutoff for {counts["publication_date_within_2022_to_cutoff_verified"]} records. Incomplete dates retain their reported precision rather than an invented day. Abstract text is available from the verified metadata for {abstract_count} records. These checks support bibliographic tracing, not final study inclusion or journal-quality certification.')
    table(['Publication year','Verified candidate records'],[[year, years[year]] for year in sorted(years)],[3.8,2.7])
    table(['Acquisition measure','Current count'],[
        ['Candidate DOI records',len(records)],
        ['Readable PDFs available',counts['readable_pdf_files_verified']],
        ['PDFs with DOI found on first two pages',counts['pdf_with_doi_on_first_two_pages']],
        ['Full text reviews completed by the retrieval script',0],
        ['Final included studies',0],
    ],[5.0,1.5])
    paragraph('The initial automated relevance audit found only 18 direct cloud/audit signals among the original 100 candidates. That is a keyword-based triage result, not a completed screening decision. Some remaining records concern adjacent network or IoT security, and some may ultimately be excluded. Expanding the register rather than forcing all 100 into the synthesis protects the distinction between a collection target and scientific eligibility.')
    paragraph('The candidate register in Appendix A preserves bibliographic identity and full-text availability. It is not the final consolidated summary of 100 fully reviewed papers requested by the guide. That deliverable requires full-text methods, findings and limitations for each included study. NotebookLM and SciSpace use has not been performed or claimed in this preparation step.')
    heading('6 Research positioning')
    paragraph('The experimental manuscript provides a concrete example of why temporal evidence matters. Its legacy controlled benchmark reports fewer false alarms after some updates, but uses the current adaptation batch for scoring. A separately versioned native-time protocol instead scores before updating and exposes a recall versus false-alarm trade-off. This example motivates the review framework. It is the author\'s own experiment, not a finding extracted from the candidate literature.')
    paragraph('The intended synthesis will ask whether explanation changes provide information beyond input drift, whether an update gate can distinguish routine changes from security-relevant shifts, and whether that distinction remains useful under transfer. It will also examine the risk of using predicted-normal records during retraining. A model can assign a normal prediction to an attack, so filtering by its own predictions cannot guarantee clean adaptation.')
    heading('7 Limitations and completion requirements')
    paragraph('The present draft has an incomplete full-text collection and no completed included-study appraisal. The original retrieval selected a fixed number of metadata hits and omitted relevant core papers that later citation follow-up recovered. These limitations preclude a claim of exhaustive coverage. Current date and language eligibility, venue legitimacy and PDF-to-article matching still require review at the study level.')
    paragraph('Before submission, the review needs a complete search update, recorded screening and exclusion decisions, at least the required number of eligible journal papers, complete extraction, appraisal and a synthesis linked to those included papers. The final PRISMA flow must use actual records at each stage. The checklist location column must refer to real manuscript sections or final page numbers. These counts and locations must not be fabricated to satisfy a deadline.')
    heading('8 Conclusion')
    paragraph('An attribution-guided update claim requires evidence about both model behaviour and subsequent detection quality. The proposed review separates these claims and organises them by data, label access, chronology and transfer. Bibliographic verification and lawful retrieval have advanced the evidence collection, but the final systematic synthesis remains unfinished. No claim about the comparative effectiveness of all candidate methods follows from metadata alone.')
    heading('Declarations')
    paragraph('The review protocol is unregistered. Generative AI assisted metadata auditing, candidate organisation and drafting. No independent reviewer agreement statistics or completed risk-of-bias ratings are reported. Funding and competing-interest declarations require author confirmation. The project number remains unspecified. Code and evidence paths are recorded in the local research package; no new public release or journal submission is claimed.')
    heading('References')
    paragraph('[1] Page MJ, McKenzie JE, Bossuyt PM, et al. The PRISMA 2020 statement: an updated guideline for reporting systematic reviews. BMJ. 2021;372:n71. https://doi.org/10.1136/bmj.n71. This methodological reference is outside the 2022 to 2026 candidate count.')
    paragraph('Candidate article references appear in Appendix A. Their presence in the register does not indicate final inclusion.')
    doc.add_page_break()
    heading('Appendix A Candidate journal article register')
    paragraph('Records P001 to P100 preserve the original candidate identifiers. CORE records are directed follow-up additions. Availability indicates a readable PDF, not completed full-text review. All records await final eligibility and critical appraisal.')
    for row in records:
        authors=row.get('verified_authors','').split('; ')
        author_text=', '.join(authors[:3])+(' et al.' if len(authors)>3 else '')
        title=re.sub(r'<[^>]+>','',html.unescape(row['verified_title'])).replace('\ufffd',' ')
        citation=html.unescape(f"{row['paper_id']}  {author_text}. {title}. {row['verified_journal']}. {row['verified_first_publication_date']}. https://doi.org/{row['doi']}")
        evidence=('Readable PDF' if row['pdf_parse_verified'] else 'Full text not retrieved')
        evidence+='\nAbstract metadata: '+('available' if row.get('verified_abstract') else 'not available')
        if row['pdf_parse_verified']:
            evidence+='\nDOI on opening pages: '+('yes' if row.get('doi_found_on_first_two_pages') else 'not found')
        p=doc.add_paragraph(citation)
        p.paragraph_format.keep_together=True
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(2)
        p.paragraph_format.line_spacing=1.05
        for run in p.runs:run.font.size=Pt(10.5)
        p=doc.add_paragraph(evidence.replace('\n','. ')+'.')
        p.paragraph_format.keep_together=True
        p.paragraph_format.space_after=Pt(8)
        p.paragraph_format.line_spacing=1.0
        for run in p.runs:run.font.size=Pt(10)
        blocks.extend([citation,evidence.replace('\n','. ')+'.'])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    (ROOT/'docs/review_paper_working_draft.md').write_text('\n\n'.join(blocks)+'\n',encoding='utf-8')
    # Preserve exact checklist wording and item identifiers from the supplied file.
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    checklist_path=Path('D:/Downloads/PRISMA_2020_checklist (1).docx')
    with ZipFile(checklist_path) as archive:
        checklist=ET.fromstring(archive.read('word/document.xml'))
    entries=[]
    for tr in checklist.findall('.//w:tr',ns):
        cells=[' '.join(''.join(t.text or '' for t in p.findall('.//w:t',ns)) for p in tc.findall('.//w:p',ns)).strip() for tc in tr.findall('w:tc',ns)]
        if len(cells)>=3 and re.fullmatch(r'\d+[a-f]?',cells[1]):
            entries.append((cells[1],cells[2]))
    available={'1':'Title identifies an in-preparation systematic review','2':'Abstract reports acquisition status and limitations',
        '3':'Section 1','4':'Section 2','5':'Section 3; final eligibility pending','6':'Section 3; search update pending',
        '7':'Archived initial queries; complete database-specific exports pending',
        '8':'Section 3; independent human screening not completed','9':'Section 3; extraction pending',
        '10a':'Section 3','10b':'Section 3','11':'Section 4; actual ratings pending','12':'Section 4; no pooled effect estimate',
        '13a':'Section 4; grouping proposed','13b':'Section 3; missing data left unreported',
        '13c':'Section 5; acquisition tables only','13d':'Section 4; narrative synthesis planned',
        '13e':'Pending included-study synthesis','13f':'Pending included-study sensitivity analysis',
        '14':'Pending','15':'Pending','16a':'Section 5; no final included-study flow',
        '16b':'Pending full-text exclusions','17':'Appendix A is candidates, not included-study characteristics',
        '18':'Pending','19':'Pending','20a':'Pending','20b':'Not performed','20c':'Not performed','20d':'Not performed',
        '21':'Pending','22':'Pending','23a':'Sections 6 and 8; no final synthesis','23b':'Section 7',
        '23c':'Section 7','23d':'Sections 6 and 8','24a':'Declarations: unregistered',
        '24b':'Local PRISMA_protocol.md','24c':'Directed follow-up recorded in Section 3',
        '25':'Declarations; author confirmation pending','26':'Declarations; author confirmation pending',
        '27':'Declarations and local audit register; public release unverified'}
    status=['# PRISMA completion map', '', 'Mapped from the supplied PRISMA checklist. Locations are sections, not unverified page numbers. No item is marked completed merely because its method is proposed.', '', '| Item | Reporting requirement | Current location or gap |','|---|---|---|']
    status += [f'| {number} | {text.replace("|","/")} | {available.get(number,"Pending")} |' for number,text in entries]
    (AUDIT/'PRISMA_completion_map.md').write_text('\n'.join(status)+'\n',encoding='utf-8')
    print(json.dumps({'output':str(OUT),'candidate_records':len(records),'body_words':len(' '.join(blocks).split()),'checklist_subitems':len(entries)},indent=2))


if __name__=='__main__':
    main()
