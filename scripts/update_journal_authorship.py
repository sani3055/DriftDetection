"""Preserve the existing manuscript while updating the agreed author order."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt

root = Path(__file__).resolve().parents[1]
source = root / 'outputs/SHAP_CloudTrail_Journal_Paper_v4.docx'
output = root / 'outputs/SHAP_CloudTrail_Journal_Paper_v4_Two_Authors.docx'
if output.exists():
    raise FileExistsError(f'Preserve existing output: {output}')
doc = Document(source)
authors = [p for p in doc.paragraphs if p.text.strip() == 'Sanidhya Bhandari']
assert len(authors) == 1
authors[0].runs[0].text = 'Sanidhya Bhandari, Badrinath N'
affiliation = [p for p in doc.paragraphs if p.text.startswith('Department of IoT, School of Computer Science')]
assert len(affiliation) == 1
element = OxmlElement('w:p')
affiliation[0]._p.addnext(element)
corresponding = Paragraph(element, affiliation[0]._parent)
corresponding.style = affiliation[0].style
corresponding.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = corresponding.add_run('Corresponding author: Badrinath N')
run.font.name = 'Times New Roman'
run.font.size = Pt(10)

ack = [p for p in doc.paragraphs if p.text.strip() in {'Acknowledgements', 'I thank Dr Badrinath N for project guidance.'}]
assert len(ack) == 2
for paragraph in ack:
    paragraph._element.getparent().remove(paragraph._element)
disclosure = [p for p in doc.paragraphs if 'The author must verify the manuscript' in p.text]
assert len(disclosure) == 1
for run in disclosure[0].runs:
    run.text = run.text.replace('The author must verify the manuscript', 'The authors must verify the manuscript')
doc.core_properties.author = 'Sanidhya Bhandari; Badrinath N'
doc.save(output)
print(output)
