"""Build the source-linked review documents from checked evidence records."""
from collections import Counter
import csv
import hashlib
import html
import json
from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT/'20106UG03'
REC = OUT/'Review_Records'
PAPERS = OUT/'Downloaded_Papers_20106UG03'
PAPERS.mkdir(exist_ok=True)
ROWS = json.loads((REC/'selected_evidence_collection.json').read_text(encoding='utf-8'))
SUMMARIES = json.loads((REC/'source_checked_summaries.json').read_text(encoding='utf-8'))
BY_ID = {r['record_id']: r for r in ROWS}
INITIAL = {
 'S0196': 'R001_CORE03_b65937b4e64ace3787dd.pdf',
 'CORE05': 'R002_CORE05_c104141fd5c3ba61e552.pdf',
 'CORE07': 'R003_CORE07_Adediran_2026_CloudTrail_RAG.pdf',
}


def norm(s):
    for _ in range(3):
        s=html.unescape(s)
    return re.sub(r'\s+', ' ', s.replace('\ufffe','').replace('\ufffd','')).strip()


def setup():
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.72)
    sec.left_margin=sec.right_margin=Inches(.8)
    sec.header_distance=sec.footer_distance=Inches(.32)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption'):
        style=d.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11)
        style.paragraph_format.space_after=Pt(6)
        style.paragraph_format.line_spacing=1.08
        if style.element.rPr is not None:
            for color in style.element.rPr.findall(qn('w:color')):
                for attr in ('themeColor','themeTint','themeShade'):
                    color.attrib.pop(qn('w:'+attr),None)
    for name,size in [('Title',22),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        d.styles[name].font.size=Pt(size)
        d.styles[name].font.bold=True
        d.styles[name].paragraph_format.space_before=Pt(10)
        d.styles[name].paragraph_format.keep_with_next=True
    for style in d.styles:
        for borders in list(style.element.iter(qn('w:pBdr'))):
            borders.getparent().remove(borders)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after=Pt(0)
    r=footer.add_run('Project 20106UG03  |  ');r.font.size=Pt(9)
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    d.core_properties.author='Sanidhya Bhandari'
    d.core_properties.subject='Explainable drift diagnosis and adaptive security monitoring'
    return d


def add_link(p,label,url):
    rel=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)


def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
    props.append(borders)
    for i,h in enumerate(headers):t.rows[0].cells[i].text=str(h)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        for c,v in zip(t.add_row().cells,values):c.text=str(v)
    for i,row in enumerate(t.rows):
        for j,c in enumerate(row.cells):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for side in ('top','left','bottom','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');pr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                if j==len(headers)-1 and len(str(c.text))<10:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(10);r.bold=(i==0)
    d.add_paragraph().paragraph_format.space_after=Pt(3)
    return t


def ref(r):
    b=r['verified_bibliography']
    authors=', '.join(b['authors']) or r.get('authors','')
    parts=[f"{authors}. {b['title'].rstrip('.')}. {b['journal']}.",str(b['date_parts'][0])]
    if b.get('volume'):parts.append('vol. '+b['volume'])
    if b.get('issue'):parts.append('no. '+b['issue'])
    if b.get('page'):parts.append('pp. '+b['page'])
    elif b.get('article_number'):parts.append('article '+b['article_number'])
    return norm(' '.join(parts)+'.')


def cite(s):
    def repl(m):
        ids=m.group(1).split(';@')
        return '['+', '.join(str(int(BY_ID[x]['review_id'][1:])) for x in ids)+']'
    return re.sub(r'\[@([^\]]+)\]',repl,s)


def evidence_label(r):
    if r.get('xml_sections'):
        return 'Article abstract and structured full text retained in source card '+r['record_id']+'.'
    return 'PDF abstract and page-linked extracts retained in source card '+r['record_id']+'.'


def copy_and_register():
    register=[]
    for r in ROWS:
        source=ROOT/r['pdf_path']
        name=INITIAL.get(r['record_id'],r['review_id']+'_'+re.sub(r'[^A-Za-z0-9]+','_',r['doi'])+'.pdf')
        dest=PAPERS/name
        digest=hashlib.sha256(source.read_bytes()).hexdigest()
        if dest.exists():
            # Initial notebook versions may have a different byte layout. Do not overwrite.
            existing=hashlib.sha256(dest.read_bytes()).hexdigest()
            if existing!=digest:
                dest=PAPERS/(r['review_id']+'_Verified_'+source.name)
        if not dest.exists():shutil.copy2(source,dest)
        pdf=pdfium.PdfDocument(str(dest))
        opening=norm(' '.join(pdf[i].get_textpage().get_text_range() for i in range(min(3,len(pdf)))))
        doi_match=re.sub(r'\s','',r['doi'].lower()) in re.sub(r'\s','',opening.lower())
        rights=r.get('archive_licence') or 'Check original PDF licence or repository rights; no public redistribution assumed'
        if not r.get('archive_licence'):
            full=' '.join(pdf[i].get_textpage().get_text_range() for i in range(len(pdf)))
            if re.search('Creative Commons',full,re.I):rights='Creative Commons statement in PDF; retain its exact attribution and restrictions'
        register.append({'review_id':r['review_id'],'source_id':r['record_id'],'doi':r['doi'],
             'title':r['verified_bibliography']['title'],'authors':'; '.join(r['verified_bibliography']['authors']),
             'journal':r['verified_bibliography']['journal'],'bibliographic_year':r['verified_bibliography']['date_parts'][0],
             'theme':SUMMARIES[r['record_id']]['theme'],'file':str(dest.relative_to(OUT)),
             'sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),'pages':len(pdf),
             'doi_found_on_opening_pages':doi_match,'rights':rights,
             'summary_complete':True,'charting':'AI assisted abstract and targeted full-text charting',
             'independent_human_appraisal':False,'final_systematic_review_inclusion':False})
    if len({x['doi'] for x in register})!=len(register):raise ValueError('Duplicate DOI')
    if len({x['sha256'] for x in register})!=len(register):raise ValueError('Duplicate PDF file')
    (REC/'Paper_Register.json').write_text(json.dumps(register,ensure_ascii=False,indent=2),encoding='utf-8')
    with (OUT/'Paper_Register.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.DictWriter(f,fieldnames=list(register[0]));w.writeheader();w.writerows(register)
    return register


def consolidated(register):
    d=setup();d.add_paragraph('Consolidated Journal Paper Summaries','Title')
    d.add_paragraph('Project 20106UG03',style='Subtitle')
    d.add_paragraph('Explainable drift diagnosis and adaptive anomaly detection in cloud and security logs')
    d.add_paragraph(f'This collection summarises {len(ROWS)} primary journal articles within 2022 to 31 October 2026. Each study is linked to a DOI and retained PDF. Reported methods are separated from reviewer interpretation so the collection can support the guide\'s review-paper assessment and further source checking.')
    d.add_paragraph('The summaries use AI-assisted abstract and targeted full-text charting. They are not independent human critical appraisals. Bibliographic verification and possession of a full text do not establish study quality or complete PRISMA screening.')
    counts=Counter(x['theme'] for x in SUMMARIES.values())
    table(d,['Dominant methodological theme','Papers'],sorted(counts.items()),[5.6,1.25])
    d.add_paragraph('How to use the collection','Heading 1')
    d.add_paragraph('Study numbers match the review draft and Paper Register. The register records PDF filenames, hashes and rights information. Page-linked source extracts or structured full text are retained in Review Records. Reviewer cautions identify questions to check; they do not assert that every paper has the same weakness.')
    d.add_paragraph('The search cutoff is not the end of 2026. Publication-date differences between online and issue metadata are recorded separately. No pooled ranking of incompatible classification, drift and update metrics is presented.')
    d.add_page_break()
    for i,(r,reg) in enumerate(zip(ROWS,register)):
        if i and i%2==0:d.add_page_break()
        start=len(d.paragraphs)
        d.add_paragraph('Study '+r['review_id'],'Heading 2')
        p=d.add_paragraph();p.add_run(norm(r['verified_bibliography']['title'])).bold=True
        p=d.add_paragraph(ref(r));p.paragraph_format.space_after=Pt(3)
        for run in p.runs:run.font.size=Pt(9.5)
        p=d.add_paragraph();add_link(p,'DOI '+r['doi'],'https://doi.org/'+r['doi'])
        p.paragraph_format.space_after=Pt(3)
        d.add_paragraph('Theme: '+SUMMARIES[r['record_id']]['theme'])
        d.add_paragraph(SUMMARIES[r['record_id']]['summary'])
        p=d.add_paragraph();p.add_run('Reviewer interpretation. ').bold=True
        p.add_run(SUMMARIES[r['record_id']]['interpretation'])
        p=d.add_paragraph(evidence_label(r));p.paragraph_format.space_after=Pt(8)
        for run in p.runs:run.font.size=Pt(9)
        # Keep one study together where possible, without fixing heights.
        for p in d.paragraphs[start:-1]:p.paragraph_format.keep_with_next=True
    target=OUT/'20106UG03_Consolidated_Journal_Paper_Summaries.docx'
    d.save(target)
    return target


def manuscript():
    d=setup()
    source=(OUT/'Review_Paper_Draft.md').read_text(encoding='utf-8')
    for line in source.splitlines():
        line=line.strip()
        if not line:continue
        if line.startswith('# '):d.add_paragraph(line[2:],'Title')
        elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 1')
        elif line.startswith('### '):d.add_paragraph(line[4:],'Heading 2')
        elif line.startswith('The Word version expands'):
            for r in ROWS:
                n=int(r['review_id'][1:]);p=d.add_paragraph('['+str(n)+'] '+ref(r)+' ')
                add_link(p,'https://doi.org/'+r['doi'],'https://doi.org/'+r['doi'])
                p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.first_line_indent=Inches(-.25)
                for run in p.runs:run.font.size=Pt(10)
            d.add_paragraph('[102] Page MJ, McKenzie JE, Bossuyt PM, et al. The PRISMA 2020 statement: an updated guideline for reporting systematic reviews. BMJ. 2021;372:n71. https://doi.org/10.1136/bmj.n71')
        else:d.add_paragraph(cite(line))
    target=OUT/'20106UG03_Review_Paper_Guide_Draft.docx';d.save(target)
    return target


def main():
    if len(ROWS)<100:raise ValueError('Fewer than 100 verified study records')
    assert {r['record_id'] for r in ROWS}==set(SUMMARIES)
    register=copy_and_register()
    docs=[consolidated(register),manuscript()]
    result={'project_number':'20106UG03','charting_date':'2026-10-31',
       'journal_article_records':len(ROWS),'distinct_dois':len({r['doi'] for r in register}),
       'distinct_pdf_hashes':len({r['sha256'] for r in register}),'completed_summaries':len(SUMMARIES),
       'opening_doi_not_found':[r['review_id'] for r in register if not r['doi_found_on_opening_pages']],
       'independent_human_critical_appraisals':0,'completed_exhaustive_systematic_review':False,
       'journal_submitted':False,'deliverables':[str(x.relative_to(ROOT)) for x in docs]}
    (REC/'package_validation.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    print(json.dumps(result,indent=2))


if __name__=='__main__':main()
