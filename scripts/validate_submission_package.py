"""Validate the generated research and supervisor-submission package.

This checks artifact presence and basic readability only. It does not replace human PRISMA
screening, full-text extraction, journal verification, or supervisor approval.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from pptx import Presentation


ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / "Project_XX_SHAP_CloudTrail_Drift"
REVIEW = ROOT / "review_2022_2026"


def require(path: Path, missing: list[str]) -> None:
    if not path.is_file():
        missing.append(str(path.relative_to(ROOT)))


def main() -> None:
    missing: list[str] = []
    required = [
        ROOT / "IMPLEMENTATION_PLAN.md",
        ROOT / "requirements.txt",
        ROOT / "docs" / "research_gap.md",
        ROOT / "docs" / "literature_review.md",
        ROOT / "docs" / "external_validation.md",
        ROOT / "docs" / "final_validation.md",
        ROOT / "review_2022_2026" / "evidence_matrix_100.csv",
        ROOT / "review_2022_2026" / "automated_quality_audit_100.csv",
        ROOT / "review_2022_2026" / "PRISMA_protocol.md",
        ROOT / "data" / "external_validation" / "external_temporal_diagnosis_metrics.json",
        ROOT / "data" / "drift" / "coupling_drift.csv",
        ROOT / "data" / "robustness" / "cost_sensitivity_results.csv",
        ROOT / "data" / "robustness" / "cost_sensitivity_winners.csv",
        PACKAGE / "Final_Consolidated_100_Paper_Summary.docx",
        PACKAGE / "Project_Report_Draft.docx",
        PACKAGE / "Capstone_IDP_Progress_Presentation.pptx",
    ]
    for path in required:
        require(path, missing)

    matrix_rows = list(csv.DictReader((REVIEW / "evidence_matrix_100.csv").open(encoding="utf-8-sig", newline="")))
    audit_rows = list(csv.DictReader((REVIEW / "automated_quality_audit_100.csv").open(encoding="utf-8-sig", newline="")))
    pdf_count = len(list((PACKAGE / "Downloaded_Papers").glob("*.pdf")))
    matrix_pdf_names = {
        Path(row["download_path"]).name
        for row in matrix_rows
        if row.get("download_path")
    }
    package_pdf_names = {path.name for path in (PACKAGE / "Downloaded_Papers").glob("*.pdf")}
    matrix_linked_pdf_count = len(package_pdf_names & matrix_pdf_names)
    supplementary_pdf_count = len(package_pdf_names - matrix_pdf_names)

    if len(matrix_rows) != 100:
        missing.append(f"evidence_matrix_100.csv has {len(matrix_rows)} rows, expected 100")
    if len(audit_rows) != 100:
        missing.append(f"automated_quality_audit_100.csv has {len(audit_rows)} rows, expected 100")

    readable = {"docx": False, "pptx": False}
    if (PACKAGE / "Final_Consolidated_100_Paper_Summary.docx").is_file():
        readable["docx"] = bool(Document(str(PACKAGE / "Final_Consolidated_100_Paper_Summary.docx")).paragraphs)
    if (PACKAGE / "Capstone_IDP_Progress_Presentation.pptx").is_file():
        readable["pptx"] = len(Presentation(str(PACKAGE / "Capstone_IDP_Progress_Presentation.pptx")).slides) == 8

    if not readable["docx"]:
        missing.append("consolidated review DOCX is not readable")
    if not readable["pptx"]:
        missing.append("presentation is not a readable 8-slide PPTX")

    metrics_path = ROOT / "data" / "external_validation" / "external_temporal_diagnosis_metrics.json"
    metrics = json.loads(metrics_path.read_text(encoding="utf-8")) if metrics_path.is_file() else {}
    temporal_metrics_present = all(metrics.get(policy) for policy in ("dual_signal", "coupled_signal"))
    if not temporal_metrics_present:
        missing.append("external temporal diagnosis metrics are missing or malformed")
    summary = {
        "matrix_records": len(matrix_rows),
        "audit_records": len(audit_rows),
        "package_pdf_count": pdf_count,
        "matrix_linked_pdf_count": matrix_linked_pdf_count,
        "supplementary_unlinked_pdf_count": supplementary_pdf_count,
        "package_docx_readable": readable["docx"],
        "package_pptx_readable": readable["pptx"],
        "external_temporal_metrics_present": temporal_metrics_present,
        "human_prisma_review_required": True,
        "valid": not missing,
        "missing_or_invalid": missing,
    }
    print(json.dumps(summary, indent=2))
    if missing:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
